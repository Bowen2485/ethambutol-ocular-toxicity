# -*- coding: utf-8 -*-
"""Add the figure placeholder notes to the staged Word file, then copy it into place.

The submission copy may be open in a word processor, which blocks the write; the staged file is
produced first and copied over when the handle is free.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
STAGE = ROOT / ".tmp" / "Ethambutol_TADS_manuscript.docx"
OUT = ROOT / "submission" / "tads" / "Ethambutol_TADS_manuscript.docx"

doc = Document(STAGE)
placed = {1: False, 2: False, 3: False, 4: False}

for p in list(doc.paragraphs):
    text = p.text
    if text.strip().startswith("Figure ") and "." in text[:12]:
        continue
    if text.strip().startswith("[Figure"):
        continue
    for n in (1, 2, 3, 4):
        if placed[n] or f"Figure {n}" not in text:
            continue
        new_p = OxmlElement("w:p")
        p._p.addnext(new_p)
        para = Paragraph(new_p, p._parent)
        run = para.add_run(f"[Figure {n} near here]")
        run.italic = True
        placed[n] = True
        break

doc.save(STAGE)
print("placeholders placed:", placed)
try:
    shutil.copy2(STAGE, OUT)
    print(f"copied to {OUT.relative_to(ROOT)}")
except PermissionError:
    print("target still open in another program; left staged")
