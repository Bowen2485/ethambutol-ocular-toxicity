# -*- coding: utf-8 -*-
"""Cover letter for the TADS submission.

A cover letter is mandatory and has to say why the manuscript suits the journal, so it is written
to the editor rather than filled with the abstract.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "tads" / "Cover_letter.docx"

TITLE = ("Optic-nerve-predominant adverse-event reporting with ethambutol: a multi-database "
         "pharmacovigilance study with exploratory transcriptomic analysis")

PARAS = [
    ("Dear Editors,", False),
    (f'We submit "{TITLE}" for consideration as an Original Article in Therapeutic Advances in '
     "Drug Safety.", False),
    ("Ethambutol optic neuropathy is well known, and several pharmacovigilance papers have already "
     "shown that the drug is reported with optic-nerve events. That is not the question we set out "
     "to answer. Our question is whether the reporting signal is concentrated in the optic nerve, "
     "or whether ethambutol simply generates more reports of every ocular kind. That distinction "
     "matters to a reader of this journal, because it determines whether a safety signal is "
     "informative about the drug or about reporting behaviour.", False),
    ("To answer it we took MedDRA's own ocular queries rather than a list assembled from the "
     "literature: all eleven narrow ocular Standardised MedDRA Queries were compared within the "
     "same de-duplicated FAERS cohort, with the optic-nerve query as one of eleven rather than as "
     "the only outcome examined. That internal comparison is the design feature we believe is "
     "novel, and it is the reason the paper can say something a single-outcome study cannot. We "
     "then asked what the pattern looks like elsewhere: JADER supports the optic-nerve signal at "
     "term level, and Canada Vigilance shows a broader, string-matched optic-related signal that we "
     "report as exploratory rather than as an equivalent replication.", False),
    ("The conclusions are bounded deliberately. Every estimate is a reporting odds ratio from "
     "spontaneous reports: it is a measure of disproportionality, not of incidence, risk or "
     "causation. The transcriptomic component is presented as an exploratory reanalysis of a "
     "public dataset at concentrations far above clinical exposure, and we say plainly where its "
     "predictions were not reproduced.", False),
    ("The manuscript follows the READUS-PV recommendations for pharmacovigilance studies and is "
     "accompanied by the completed checklist. It is not published, not under consideration "
     "elsewhere, and all authors have read and approved the submitted version. The authors declare "
     "no competing interests, and the work received no specific funding. Data and code are "
     "available at https://github.com/Bowen2485/ethambutol-ocular-toxicity under the release tag "
     "v2.0-tads-2026-10-21.", False),
    ("We hope the manuscript is suitable for the journal and look forward to your assessment.", False),
    ("Yours sincerely,", False),
    ("Bowen Shi, on behalf of all authors", False),
    ("Eye Hospital, China Academy of Chinese Medical Sciences, Beijing, China", False),
    ("20200121047@bucm.edu.cn", False),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.15

for text, bold in PARAS:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold

doc.save(OUT)
print(f"written {OUT.relative_to(ROOT)}")
