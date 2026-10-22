# -*- coding: utf-8 -*-
"""Produce a standalone Figure legends file for submission.

The legends already appear in the manuscript before the tables, but journals ask for them as a
separate upload too, and the review noted that the figure pages carry only the images.  Each
legend is extracted from the assembled manuscript so the two cannot drift.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
text = (ROOT / "manuscript" / "TADS_manuscript.md").read_text(encoding="utf-8")
start = text.find("## Figure legends")
end = text.find("## Tables", start)
block = text[start:end]

legends = re.findall(r"\*\*(Figure \d+\.)\*\*\s*(.+?)(?=\n\*\*Figure|\n## |\Z)", block, re.S)
assert len(legends) == 4, f"expected 4 legends, found {len(legends)}"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

head = doc.add_paragraph()
run = head.add_run("Figure legends")
run.bold = True
run.font.size = Pt(13)

for label, body in legends:
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(" ".join(body.split()))
    p.paragraph_format.space_after = Pt(6)

out = ROOT / "submission" / "tads" / "Figure_legends.docx"
doc.save(out)
print(f"written {out.relative_to(ROOT)} ({len(legends)} legends)")
for label, body in legends:
    print(f"  {label} {' '.join(body.split())[:80]}...")
