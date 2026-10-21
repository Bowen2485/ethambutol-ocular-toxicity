# -*- coding: utf-8 -*-
"""Render the READUS-PV checklist as a standalone Word file for submission.

The Methods say the completed checklist is supplied; a Markdown file in docs/ is not a submission
document, so the table is converted to a .docx that can be uploaded as a supplementary file.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "READUS-PV_checklist.md"
OUT = ROOT / "submission" / "tads" / "READUS-PV_checklist.docx"

lines = SRC.read_text(encoding="utf-8").splitlines()
rows: list[list[str]] = []
for ln in lines:
    if ln.startswith("|") and not re.match(r"^\|[\s:\-|]+\|$", ln):
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if cells and cells[0].lower() != "#":
            rows.append(cells)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(9)

title = doc.add_paragraph()
run = title.add_run("READUS-PV checklist — completed")
run.bold = True
run.font.size = Pt(13)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT

sub = doc.add_paragraph()
sub.add_run(
    "Manuscript: Optic-nerve-predominant adverse-event reporting with ethambutol: a multi-database "
    "pharmacovigilance study with exploratory transcriptomic analysis.\n"
    "Checklist: READUS-PV (Fusaroli M, et al. Drug Saf. 2024;47(6):575-584), 32 recommendations in "
    "14 items. Section numbers refer to the manuscript; the numbers refer to the published checklist."
).font.size = Pt(9)

table = doc.add_table(rows=0, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
header = ["#", "Recommendation (abridged)", "Where reported", "How it is addressed"]
hdr = table.add_row().cells
for cell, text in zip(hdr, header):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.size = Pt(9)

for row in rows:
    if len(row) < 4:
        continue
    cells = table.add_row().cells
    for cell, text in zip(cells, row[:4]):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.size = Pt(8.5)

widths = [Pt(24), Pt(150), Pt(80), Pt(210)]
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = width

doc.add_paragraph()
note = doc.add_paragraph()
note.add_run(
    "Items not met, stated rather than implied: no case-by-case review or causality assessment was "
    "performed (items 7d and 10); the analysis was not prospectively registered, so the transcriptomic "
    "predictions are labelled post hoc and the query and control sets are described as nominated in "
    "advance of the analysis but after the data were available (item 14d)."
).font.size = Pt(8.5)

doc.save(OUT)
print(f"written {OUT.relative_to(ROOT)} ({len(rows)} rows)")
