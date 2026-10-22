# -*- coding: utf-8 -*-
"""Repair the Word manuscript: leaked markdown and heading bullets.

Three faults are visible when the file is opened:

    "### 2.4 Exposure" was printed as body text, because only "#" and "##" were converted
    "---" separators were printed as literal dashes
    headings carried a bullet in the word processor

The first two are converter bugs. The third comes from number formatting the application applies
to heading styles, so each heading now states explicitly that it is unnumbered.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
p = ROOT / "src" / "build_manuscript_docx.py"
s = p.read_text(encoding="utf-8")

#: explicit "no numbering" on every heading
helper = '''

def unnumbered(par) -> None:
    """Force numId=0 so the word processor does not draw a bullet before a heading."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = par._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)

'''
s = s.replace("def add_runs(par, text: str) -> None:", helper.strip("\n") + "\n\ndef add_runs(par, text: str) -> None:", 1)

#: heading levels 1-4, all unnumbered, and horizontal rules dropped
old = '''    if s.startswith("# "):
        h = doc.add_heading(level=1)
        add_runs(h, s[2:])
        i += 1
        continue
    if s.startswith("## "):
        h = doc.add_heading(level=2)
        add_runs(h, s[3:])
        i += 1
        continue'''
new = '''    if re.fullmatch(r"-{3,}|\\*{3,}", s):
        i += 1
        continue
    m = re.match(r"^(#{1,4})\\s+(.*)$", s)
    if m:
        level = len(m.group(1))
        h = doc.add_heading(level=level)
        add_runs(h, m.group(2))
        unnumbered(h)
        i += 1
        continue'''
assert old in s, "heading branch not found"
s = s.replace(old, new, 1)

#: the loop that gathers a paragraph must also stop at any heading level and at rules
old = 'while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "|")):'
new = 'while (i < len(lines) and lines[i].strip()\n           and not lines[i].strip().startswith(("#", "|"))\n           and not re.fullmatch(r"-{3,}|\\*{3,}", lines[i].strip())):'
assert old in s, "paragraph gather not found"
s = s.replace(old, new, 1)

p.write_text(s, encoding="utf-8")
print("builder patched: headings 1-4, rules skipped, headings unnumbered")
